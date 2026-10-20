import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_directory_deck(output_path="NPF_National_Command_Directory.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Executive Theme Colors
    BG_DARK = RGBColor(10, 22, 40)        # #0A1628 Deep Midnight
    BG_CARD = RGBColor(16, 34, 60)        # #10223C Dark Card Navy
    BG_CARD_BORDER = RGBColor(30, 58, 95) # #1E3A5F Border
    EMERALD = RGBColor(0, 196, 106)       # #00C46A Glowing Police Green
    CYAN = RGBColor(13, 204, 176)         # #0DCCB0 Cyan
    GOLD = RGBColor(245, 158, 11)         # #F59E0B Gold
    WHITE = RGBColor(255, 255, 255)
    MUTED = RGBColor(148, 163, 184)       # #94A3B8 Slate Gray
    LIGHT_TEXT = RGBColor(226, 232, 240)  # #E2E8F0 Soft White
    RED_ACCENT = RGBColor(239, 68, 68)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    police_logo = os.path.join(base_dir, "sitrep-app", "public", "police.png")

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()

        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = EMERALD
        accent.line.fill.background()
        return bg

    def add_header(slide, title, kicker="NIGERIA POLICE FORCE • COMMAND DIRECTORY & HIERARCHY"):
        kb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.3))
        tf_k = kb.text_frame
        tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = kicker.upper()
        p_k.font.size = Pt(10)
        p_k.font.bold = True
        p_k.font.color.rgb = EMERALD

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(10.5), Inches(0.75))
        tf_t = tb.text_frame
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

        if os.path.exists(police_logo):
            slide.shapes.add_picture(police_logo, Inches(12.0), Inches(0.4), Inches(0.65), Inches(0.75))

    def add_card(slide, left, top, width, height, title, items, badge_text=None, accent_color=EMERALD):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BG_CARD_BORDER
        card.line.width = Pt(1.2)

        hbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.06))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = accent_color
        hbar.line.fill.background()

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = WHITE
        p0.space_after = Pt(8)

        if badge_text:
            p_badge = tf.add_paragraph()
            p_badge.text = f"[{badge_text}]"
            p_badge.font.size = Pt(9.5)
            p_badge.font.bold = True
            p_badge.font.color.rgb = accent_color
            p_badge.space_after = Pt(6)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(10)
            p.font.color.rgb = MUTED
            p.space_after = Pt(4)

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)
    if os.path.exists(police_logo):
        s1.shapes.add_picture(police_logo, Inches(5.9), Inches(1.1), Inches(1.5), Inches(1.75))

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(3.1), Inches(11.33), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    p0.text = "OFFICIAL DIRECTORY & FORMATION HIERARCHY"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = EMERALD
    p0.space_after = Pt(10)

    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "NIGERIA POLICE FORCE COMMAND DIRECTORY"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Force Headquarters • 17 Zonal Commands • 37 State Commands • Area Commands • Divisions & Stations"
    p2.font.size = Pt(14)
    p2.font.color.rgb = CYAN
    p2.space_after = Pt(20)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.16), Inches(5.6), Inches(7.0), Inches(0.9))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BG_CARD
    badge.line.color.rgb = BG_CARD_BORDER
    tf_b = badge.text_frame
    pb1 = tf_b.paragraphs[0]
    pb1.alignment = PP_ALIGN.CENTER
    pb1.text = "Comprehensive Territorial Roster & Administrative Chain of Command"
    pb1.font.size = Pt(11)
    pb1.font.bold = True
    pb1.font.color.rgb = GOLD
    pb2 = tf_b.add_paragraph()
    pb2.alignment = PP_ALIGN.CENTER
    pb2.text = "Compiled from Official NPF Gazette & Public Records for Election Operations"
    pb2.font.size = Pt(9.5)
    pb2.font.color.rgb = MUTED

    # ==========================================
    # SLIDE 2: NPF COMMAND TIERS & APEX FORMATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "NPF Command Hierarchy: 5-Tier Operational Cascade")

    tiers = [
        ("Tier 1: Force Headquarters (FHQ)", "Louis Edet House, Area 11, Garki, FCT Abuja", "Led by Inspector-General of Police (IGP) assisted by 7 DIGs (Departments: Operations, Admin, Logistics, Training, Intelligence, FCID, ICT).", GOLD),
        ("Tier 2: 17 Zonal Commands (Zones 1-17)", "Strategic regional headquarters across geopolitical zones", "Headed by Assistant Inspectors-General of Police (AIG). Supervise between 2 and 3 contiguous State Police Commands.", EMERALD),
        ("Tier 3: 37 State Police Commands", "36 State Commands + Federal Capital Territory (FCT)", "Led by Commissioners of Police (CP). Exercises overall territorial security control within state borders.", CYAN),
        ("Tier 4: Area Commands (~150+ Nationwide)", "Sub-state operational hubs covering clusters of LGAs", "Commanded by Assistant Commissioners of Police (ACP) or CSP. Coordinates multiple local police divisions.", WHITE),
        ("Tier 5: Divisional Police HQs & Stations", "Primary grassroots frontline units embedded in communities", "Led by Divisional Police Officers (DPO - CSP / SP). Directly manages Police Stations, Outposts, and Police Posts.", LIGHT_TEXT)
    ]

    top_y = 1.7
    for title, subtitle, desc, col in tiers:
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_y), Inches(11.73), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BG_CARD_BORDER

        bar = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top_y), Inches(0.12), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = s2.shapes.add_textbox(Inches(1.1), Inches(top_y + 0.1), Inches(11.2), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"{title} — {subtitle}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = MUTED

        top_y += 1.05

    # ==========================================
    # SLIDE 3: ALL 17 ZONAL COMMANDS (DIRECTORY PART 1)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Directory of 17 Zonal Police Commands (Zones 1 – 17)")

    zones_left = [
        ("Zone 1", "Kano (Bompai)", "Kano, Jigawa"),
        ("Zone 2", "Lagos (Onikan)", "Lagos, Ogun"),
        ("Zone 3", "Yola", "Adamawa, Taraba"),
        ("Zone 4", "Makurdi", "Benue, Plateau, Nasarawa"),
        ("Zone 5", "Benin City", "Edo, Delta"),
        ("Zone 6", "Calabar", "Cross River, Akwa Ibom"),
        ("Zone 7", "Abuja (FCT)", "FCT Abuja, Niger State"),
        ("Zone 8", "Lokoja", "Kogi, Kwara"),
        ("Zone 9", "Umuahia", "Abia, Imo, Ebonyi")
    ]

    zones_right = [
        ("Zone 10", "Sokoto", "Sokoto, Kebbi, Zamfara"),
        ("Zone 11", "Osogbo", "Osun, Oyo"),
        ("Zone 12", "Bauchi", "Bauchi, Gombe"),
        ("Zone 13", "Ukpo / Awka", "Anambra, Enugu"),
        ("Zone 14", "Katsina", "Katsina, Kaduna"),
        ("Zone 15", "Maiduguri", "Borno, Yobe"),
        ("Zone 16", "Yenagoa", "Bayelsa, Rivers"),
        ("Zone 17", "Akure", "Ondo, Ekiti"),
        ("FHQ Direct", "Louis Edet House", "Border Patrol, Marine, Air Wing, Rail")
    ]

    def render_zone_table(slide, zones, left_in, top_in, w_in):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_in), Inches(top_in), Inches(w_in), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BG_CARD_BORDER

        tb = slide.shapes.add_textbox(Inches(left_in + 0.2), Inches(top_in + 0.2), Inches(w_in - 0.4), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = "ZONE          HEADQUARTERS         STATE COMMANDS COVERED"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = EMERALD
        p0.space_after = Pt(8)

        for z, hq, states in zones:
            p = tf.add_paragraph()
            p.text = f"{z:<12} {hq:<22} {states}"
            p.font.size = Pt(10)
            p.font.color.rgb = WHITE
            p.space_after = Pt(4)

    render_zone_table(s3, zones_left, 0.8, 1.7, 5.7)
    render_zone_table(s3, zones_right, 6.8, 1.7, 5.7)

    # ==========================================
    # SLIDE 4: DEEP DIVE: LAGOS STATE POLICE COMMAND (AREA COMMANDS A-P)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Lagos State Command: Area Commands & Key Divisions")

    areas_lagos = [
        ("Area A (Lion Building)", "Lagos Island, VI, Ikoyi", "Lion Bldg, Central PS, Victoria Island, Ikoyi, Onikan, Bar Beach, Adeniji Adele"),
        ("Area B (Apapa)", "Apapa Port, Ajegunle, Ajeromi", "Apapa Division, Ajegunle, Trinity, Layeni, Kirikiri, Tolu, Ijora Badia"),
        ("Area C (Surulere)", "Surulere, Yaba, Ebute Metta", "Surulere Div, Yaba, Bode Thomas, Denton, Iponri, Aguda, Sabo, Orile"),
        ("Area D (Mushin)", "Mushin, Isolo, Oshodi", "Mushin Division, Isolo, Ejigbo, Ajao Estate, Aswani, Alakara, Ilasamaja, Itire"),
        ("Area E (Festac)", "Festac Town, Ojo, Satellite", "Festac Division, Ojo, Okokomaiko, Satellite Town, Trade Fair, Iba, Ajangbadi"),
        ("Area F (Ikeja)", "State Capital, Alausa, Airport", "Ikeja Division, Alausa, Ojodu, Ilupeju, Shogunle, Makinde, Airport/Domestic"),
        ("Area G (Ogba)", "Ogba, Agege, Pen Cinema", "Ogba Division, Agege, Isokoko, Elere, Dopemu, Oko-Oba, Abattoir"),
        ("Area H (Ogudu)", "Ogudu, Ketu, Ojota, Kosofe", "Ogudu Division, Ketu, Alapere, Owode Onirin, Mile 12, Anthony"),
        ("Area J (Elemoro)", "Lekki, Ajah, Epe Corridor", "Ajah Division, Maroko, Ilasan, Elemoro, Ogombo, Langbasa, Epe, Akodo"),
        ("Area K (Morogbo)", "Badagry Expressway, Border", "Badagry Division, Morogbo, Ilemba Hausa, Ishashi, Ojo Border, Seme"),
        ("Area M (Idimu)", "Alimosho, Egbe, Idimu", "Idimu Division, Egbe, Ikotun, Igando, Shasha, Gowon Estate, Meiran"),
        ("Area N (Ijede)", "Ikorodu, Sagamu Road, Imota", "Ikorodu Division, Sagamu Rd, Igbogbo, Ijede, Imota, Owutu, Shagamu Interchange"),
        ("Area P (Alagbado)", "Ipaja, Ayobo, Alagbado", "Alagbado Division, Ipaja, Ayobo, Abesan Estate, Meiran, Baruwa")
    ]

    card_lagos = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.73), Inches(5.2))
    card_lagos.fill.solid()
    card_lagos.fill.fore_color.rgb = BG_CARD
    card_lagos.line.color.rgb = BG_CARD_BORDER

    tb_l = s4.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.33), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p_hdr = tf_l.paragraphs[0]
    p_hdr.text = "AREA COMMAND (HEADQUARTERS)      JURISDICTION / DISTRICT        KEY SUBORDINATE DIVISIONS & STATIONS"
    p_hdr.font.size = Pt(10)
    p_hdr.font.bold = True
    p_hdr.font.color.rgb = EMERALD
    p_hdr.space_after = Pt(6)

    for ac, jur, divs in areas_lagos:
        p = tf_l.add_paragraph()
        p.text = f"{ac:<32} {jur:<30} {divs}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = WHITE
        p.space_after = Pt(3)

    # ==========================================
    # SLIDE 5: DEEP DIVE: FCT ABUJA POLICE COMMAND
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Federal Capital Territory (FCT) Police Command")

    fct_areas = [
        ("Area Command Metro (Abuja City)", "Maitama (Nile St)", "Central Business District, Maitama, Wuse, Asokoro, Garki, Utako, Jabi, Lugbe",
         ["Central Police Station (CPS Fed Secretariat)", "Maitama Division (Nile St)", "Wuse Division (Wuse Zone 3)", "Garki Division (Garki II)", "Asokoro Division (AYA/Gowon Cres)", "Lugbe Division (Airport Rd)", "Utako Division", "Wuye & Durumi Divisions"]),
        ("Area Command Gwagwalada", "Specialist Hospital Rd", "Gwagwalada Area Council, Airport Corridor, Kuje boundary",
         ["Gwagwalada Main Division", "Gwagwalada Main Market Division", "Iddo Division (Iddo Sarki / Airport Rd)", "Zuba Division", "Tungan Maje Police Post", "Dobi Police Outpost"]),
        ("Area Command Kubwa", "Kubwa Township", "Bwari/AMAC Northwest, Dei-Dei, Gwarimpa extensions",
         ["Kubwa Phase 4 Division", "Dei-Dei Division / Building Market", "Dutse Alhaji Division", "Byazhin Division", "Gwarimpa Division (1st & 2nd Ave)", "Dawaki Outpost"]),
        ("Area Command Bwari", "Bwari Town", "Bwari Area Council, Niger/Kaduna boundary corridors",
         ["Bwari Main Divisional HQ", "Ushafa Division", "Gawu Border Post", "Gidandogo Division", "Kuduru Police Station", "Garam Outpost"]),
        ("Karu / Karshi (Eastern Axis)", "Karu / Nyanya Corridors", "Abuja-Keffi Highway border axis with Nasarawa State",
         ["Karu Divisional HQ", "Nyanya Division (Check point/Market)", "Jikwoyi Division", "Karshi Division", "Orozo Division", "Kurudu Model Station"])
    ]

    top_y = 1.7
    for name, loc, jur, divs in fct_areas:
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_y), Inches(11.73), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BG_CARD_BORDER

        bar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top_y), Inches(0.12), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CYAN
        bar.line.fill.background()

        tb = s5.shapes.add_textbox(Inches(1.1), Inches(top_y + 0.08), Inches(11.2), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p0 = tf.paragraphs[0]
        p0.text = f"{name} • HQ: {loc} ({jur})"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = GOLD

        p_divs = tf.add_paragraph()
        p_divs.text = f"Divisions: {', '.join(divs)}"
        p_divs.font.size = Pt(9.5)
        p_divs.font.color.rgb = WHITE

        top_y += 1.05

    # ==========================================
    # SLIDE 6: MAJOR STATE COMMANDS: KANO, RIVERS, KADUNA, ANAMBRA
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Strategic State Commands: Area Commands & Divisions")

    state_samples = [
        ("Kano State Command (Zone 1)", "Bompai, Kano", "Metro, Bichi, Wudil, Rano, Dala, Danbatta, Tudun Wada, Eastern Bypass",
         "Central PS, Bichi Division, Gwarzo, Wudil, Rano, Fagge, Sabon Gari, Nassarawa, Dala, Kwalli, Gwale, Mandawari"),
        ("Rivers State Command (Zone 16)", "Moscow Rd, Port Harcourt", "Port Harcourt Metro, Bori (Ogoni), Choba, Ahoada, Eleme, Oyigbo",
         "CPS Port Harcourt, Borokiri, Mile 1 Diobu, Choba Div, Rumuolumeni, Eleme, Onne, Bori, Ahoada, Omoku, Mini-Okoro"),
        ("Kaduna State Command (Zone 14)", "Bida Rd, Kaduna", "Kaduna Metro, Zaria, Kafanchan, Birnin Gwari",
         "Sabon Gari Div, Kabala Doki, Unguwan Rimi, Kawo, Tudun Wada, Zaria City, Danmagaji, Kafanchan, Samaru, Birnin Gwari"),
        ("Anambra State Command (Zone 13)", "Awka", "Awka Metro, Onitsha, Nnewi, Ihiala, Otuocha, Aguata",
         "CPS Awka, B Div Awka, CPS Onitsha, Inland Town, Fegge, Okpoko, CPS Nnewi, Otolo, Ihiala, Otuocha, Aguata, Umunze")
    ]

    top_y = 1.7
    for st_title, st_hq, area_cmds, div_samples in state_samples:
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_y), Inches(11.73), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BG_CARD_BORDER

        bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top_y), Inches(0.12), Inches(1.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = EMERALD
        bar.line.fill.background()

        tb = s6.shapes.add_textbox(Inches(1.1), Inches(top_y + 0.1), Inches(11.2), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"{st_title} • State HQ: {st_hq}"
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = GOLD

        p1 = tf.add_paragraph()
        p1.text = f"Area Commands: {area_cmds}"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = CYAN

        p2 = tf.add_paragraph()
        p2.text = f"Key Divisions: {div_samples}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = WHITE

        top_y += 1.3

    # ==========================================
    # SLIDE 7: DIVISIONAL HQ TO POLICE POST / PU MAPPING
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Grassroots Structure: Divisions, Outposts & Polling Units")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Divisional Police Headquarters",
             [
                 "Led by Divisional Police Officer (DPO) of CSP or SP rank.",
                 "Coordinates 24/7 patrol coverage across all wards in the division.",
                 "Maintains the Divisional Crime Branch (DCB) and operations room.",
                 "Serves as the primary logistics dispatch point for ballot box escort teams."
             ],
             "LOCAL COMMAND CORE", EMERALD)

    add_card(s7, Inches(4.85), Inches(1.8), Inches(3.6), Inches(4.8),
             "Police Stations & Outposts",
             [
                 "Sub-stations commanded by Station Officers (Inspectors / ASPs).",
                 "Embedded directly in residential estates, markets, and remote towns.",
                 "Operates initial complaint charge offices and rapid-intervention cells.",
                 "Provides first-responder backup within 5–15 minutes of incident call."
             ],
             "NEIGHBORHOOD FRONT LINE", CYAN)

    add_card(s7, Inches(8.9), Inches(1.8), Inches(3.6), Inches(4.8),
             "Polling Unit (PU) Deployment",
             [
                 "1 to 2 dedicated police officers deployed per Polling Unit.",
                 "Equipped with Election SitRep mobile PWA for hourly status submissions.",
                 "Direct link to Divisional DPO and Area Command Quick Reaction Force (QRF).",
                 "Custody of stamped EC8A tally sheets from voting close to ward collation center."
             ],
             "GROUND POLICE DETACHMENT", GOLD)

    # ==========================================
    # SLIDE 8: INTEGRATION WITH ELECTION SITREP
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Operational Integration with Election SitRep Platform")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "1. Multi-Tiered User Roster Mapping",
             [
                 "Platform user accounts are mapped directly to official NPF service numbers.",
                 "Strict territorial scoping: CP (State), ACP (Area Command), DPO (Division), PU Officer.",
                 "Ensures incident feeds and dispatch rights reflect real operational authority."
             ],
             "ACCESS CONTROL MATRIX", EMERALD)

    add_card(s8, Inches(6.9), Inches(1.8), Inches(5.6), Inches(2.3),
             "2. Instant Escalation from PU to Force HQ",
             [
                 "Critical SOS alerts from a Polling Unit instantly ring at Divisional HQ, Area Command, and State Command.",
                 "Situation Room display walls immediately highlight the distressed station on GIS map.",
                 "Eliminates latency inherent in verbal radio relay trees during high-stress moments."
             ],
             "AUTOMATED ALERT CASCADE", RED_ACCENT)

    add_card(s8, Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.35),
             "3. QRF Patrol Dispatch & Geospatial Tracking",
             [
                 "Area Commands and Divisional HQs track standby Quick Reaction patrol vans.",
                 "Live turn-by-turn routing to affected polling stations upon distress trigger.",
                 "Visual audit trail of team mobilization, arrival on scene, and situation de-escalation."
             ],
             "RAPID INTERVENTION ENGINE", CYAN)

    add_card(s8, Inches(6.9), Inches(4.35), Inches(5.6), Inches(2.35),
             "4. Evidentiary Audit Trail & Post-Election Report",
             [
                 "Tamper-proof hash-chained records of all reports submitted across every Division.",
                 "Official EC8A result photographs archived with officer signature and timestamp.",
                 "Court-ready evidentiary bundles exportable for election petition tribunals."
             ],
             "LEGAL & TRIBUNAL AUDIT", GOLD)

    # Save presentation
    prs.save(output_path)
    print(f"Directory presentation saved to: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "e:/projects/ElectionSitrep/NPF_National_Command_Directory.pptx"
    create_directory_deck(out)
